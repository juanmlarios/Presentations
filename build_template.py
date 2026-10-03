"""Build reusable native PowerPoint layouts from the original deck.

Run: python3 build_template.py
Requires the already-installed lxml and python-pptx packages.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import posixpath
import uuid
import zipfile

from lxml import etree as ET
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Understanding_Neurodivergence.pptx"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
PALETTE = {
    "dk1": "1C2A2E", "lt1": "FFFFFF", "dk2": "55676C", "lt2": "F4F6F7",
    "accent1": "1C565D", "accent2": "528B8E", "accent3": "ADD3D5",
    "accent4": "E4D6BA", "accent5": "0A3A41", "accent6": "E7EEF0",
    "hlink": "1C565D", "folHlink": "528B8E",
}
# Source slide, named layout, ordered editable fields (excluding footer/number).
LAYOUTS = [
    (1, "01 Cover", [
        ("Presentation type", "PRESENTATION TEMPLATE"),
        ("Presentation title", "Your presentation\ntitle"),
        ("Subtitle", "A clear subtitle or short description"),
        ("Presenter", "Presenter name • Organisation"),
    ]),
    (2, "02 Introduction + Bullets", [
        ("Section label", "GETTING STARTED"),
        ("Slide title", "Build your presentation from these layouts"),
        ("Lead statement", "Use Home → New Slide → Layout to choose a design. Click a placeholder and replace its text."),
        ("Main points", "Choose a layout that fits your message.\nKeep each slide focused on one idea.\nUse Home → Reset to restore placeholder positions.\nDuplicate these examples or insert a fresh slide from the layout gallery."),
    ]),
    (3, "03 Four Cards Across", [
        ("Section label", "FOUR KEY IDEAS"),
        ("Slide title", "Compare four ideas at a glance"),
        ("Subtitle", "Use short headings and a few lines of supporting text."),
        *[(role, text) for i in range(1, 5) for role, text in [
            (f"Card {i} heading", f"Key idea {i}"),
            (f"Card {i} description", "Explain the idea in one or two sentences. Keep each card similar in length."),
        ]],
    ]),
    (4, "04 Contrasting Panels", [
        ("Section label", "COMPARISON"),
        ("Slide title", "Show two contrasting perspectives"),
        ("Left panel heading", "Current approach"),
        ("Left panel points", "Describe the starting point.\nIdentify the main challenge.\nName an important constraint.\nExplain what needs to change."),
        ("Right panel heading", "Proposed approach"),
        ("Right panel points", "Describe the desired outcome.\nExplain the alternative.\nHighlight a practical benefit.\nIdentify the next step."),
    ]),
    (5, "05 Eight Detail Blocks", [
        ("Section label", "OVERVIEW"),
        ("Slide title", "Summarise eight related concepts"),
        *[(role, text) for i in range(1, 9) for role, text in [
            (f"Block {i} heading", f"Concept {i}"),
            (f"Block {i} description", "Add one concise sentence explaining this concept and why it matters."),
        ]],
    ]),
    (7, "06 Four Cards Grid", [
        ("Section label", "HIGHLIGHTS"),
        ("Slide title", "Organise four observations or recommendations"),
        *[(f"Card {i} text", f"Observation {i}: Describe a useful insight or recommendation in one or two sentences.") for i in range(1, 5)],
    ]),
    (9, "07 Two Columns + Takeaway", [
        ("Section label", "TWO PERSPECTIVES"),
        ("Slide title", "Explore a topic from two sides"),
        ("Left column points", "Introduce the first perspective.\nSupport it with a clear observation.\nExplain its practical implications."),
        ("Takeaway", "Takeaway: State the single most important message you want people to remember."),
        ("Right column points", "Introduce the second perspective.\nDescribe a complementary insight.\nExplain how the two views connect."),
    ]),
    (10, "08 Six Detail Blocks + Takeaway", [
        ("Section label", "FRAMEWORK"),
        ("Slide title", "Explain six parts of a framework"),
        *[(role, text) for i in range(1, 7) for role, text in [
            (f"Block {i} heading", f"Component {i}"),
            (f"Block {i} description", "Briefly describe this component and its role in the overall framework."),
        ]],
        ("Takeaway", "Takeaway: Explain how these components work together."),
    ]),
    (13, "09 Bullet List + Takeaway", [
        ("Section label", "KEY POINTS"),
        ("Slide title", "Focus on the main points"),
        ("Main points", "Start with the most important point.\nAdd supporting evidence or a practical example.\nExplain a relevant implication or limitation.\nEnd with a clear action or recommendation."),
        ("Takeaway", "Takeaway: Distil the list into one memorable sentence."),
    ]),
    (15, "10 Case Study + Discussion", [
        ("Section label", "DISCUSSION"),
        ("Slide title", "Present a case and invite discussion"),
        ("Case label", "SCENARIO"),
        ("Case narrative", "Describe a person, team, or situation relevant to your audience. Include enough context to make the scenario concrete.\n\nExplain the challenge, the constraints, and what has already been tried. End with a decision or question for the group to explore."),
        ("Discussion label", "CONSIDER"),
        ("Discussion prompts", "What stands out?\nWhat information is missing?\nWhich assumptions should we question?\nWhat alternatives could we consider?\nWhat would you recommend next?"),
    ]),
]


def q(prefix, name):
    return f"{{{NS[prefix]}}}{name}"


def xml(element):
    return ET.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def texts(shape):
    return "".join(shape.xpath(".//a:t/text()", namespaces=NS))


def replace_text(shape, value):
    """Keep paragraph/run formatting, including bullet settings and font size."""
    body = shape.find("p:txBody", NS)
    original = next(p for p in body.findall("a:p", NS) if p.find("a:r", NS) is not None)
    template = deepcopy(original)
    for p in body.findall("a:p", NS):
        body.remove(p)
    for line in value.split("\n"):
        p = deepcopy(template)
        run = deepcopy(p.find("a:r", NS))
        for child in list(p):
            if child.tag not in [q("a", "pPr")]:
                p.remove(child)
        run.find("a:t", NS).text = line
        p.append(run)
        end = deepcopy(run.find("a:rPr", NS))
        end.tag = q("a", "endParaRPr")
        p.append(end)
        # Defaults are needed when PowerPoint inserts a fresh placeholder.
        prop = p.find("a:pPr", NS)
        if prop is None:
            prop = ET.Element(q("a", "pPr"))
            p.insert(0, prop)
        default = prop.find("a:defRPr", NS)
        if default is not None:
            prop.remove(default)
        default = deepcopy(run.find("a:rPr", NS))
        default.tag = q("a", "defRPr")
        prop.append(default)
        body.append(p)
    # Carry the field's own formatting into all nine text levels.
    style = body.find("a:lstStyle", NS)
    style.clear()
    prop = body.find("a:p/a:pPr", NS)
    default_paragraph = deepcopy(prop)
    default_paragraph.tag = q("a", "defPPr")
    style.append(default_paragraph)
    for level in range(1, 10):
        item = deepcopy(prop)
        item.tag = q("a", f"lvl{level}pPr")
        if item.find("a:buChar", NS) is not None:
            item.set("marL", str(342900 + (level - 1) * 457200))
            item.set("indent", "-342900")
        style.append(item)
    bodyprop = body.find("a:bodyPr", NS)
    bodyprop.set("wrap", "square")
    for autofit in list(bodyprop):
        if ET.QName(autofit).localname in ["normAutofit", "spAutoFit", "noAutofit"]:
            bodyprop.remove(autofit)
    ET.SubElement(bodyprop, q("a", "normAutofit"))


def theme_refs(root):
    inverse = {v: k for k, v in PALETTE.items() if k not in ["hlink", "folHlink"]}
    for color in root.findall(".//a:srgbClr", NS):
        key = inverse.get(color.get("val"))
        if key:
            color.tag = q("a", "schemeClr")
            color.set("val", key)
    for font in root.findall(".//a:latin", NS):
        font.set("typeface", "+mn-lt")
        for attr in ["panose", "pitchFamily", "charset"]:
            font.attrib.pop(attr, None)
    # A real Unicode square avoids Wingdings substitution on other platforms.
    for bullet in root.findall(".//a:buChar", NS):
        bullet.set("char", "▪")
    for font in root.findall(".//a:buFont", NS):
        font.set("typeface", "Georgia")
        font.attrib.pop("charset", None)
        font.attrib.pop("pitchFamily", None)


def placeholder(shape, index, kind, name):
    nv = shape.find("p:nvSpPr/p:nvPr", NS)
    ET.SubElement(nv, q("p", "ph"), type=kind, idx=str(index), hasCustomPrompt="1")
    shape.find("p:nvSpPr/p:cNvPr", NS).set("name", name)
    locks = ET.SubElement(shape.find("p:nvSpPr/p:cNvSpPr", NS), q("a", "spLocks"), noGrp="1")
    return locks


def set_geometry(shape, x=None, y=None, w=None, h=None):
    transform = shape.find("p:spPr/a:xfrm", NS)
    for node, attr, value in [("off", "x", x), ("off", "y", y), ("ext", "cx", w), ("ext", "cy", h)]:
        if value is not None:
            transform.find(f"a:{node}", NS).set(attr, str(round(value * 914400)))


def relationships(items):
    root = ET.Element(q("rel", "Relationships"), nsmap={None: NS["rel"]})
    for rid, kind, target in items:
        ET.SubElement(root, q("rel", "Relationship"), Id=rid, Type=f"{NS['r']}/{kind}", Target=target)
    return xml(root)


def build():
    with zipfile.ZipFile(SOURCE) as source:
        data = {n: source.read(n) for n in source.namelist()}
    parts = {n: v for n, v in data.items() if not n.startswith(("ppt/slides/", "ppt/slideLayouts/", "ppt/notesSlides/", "ppt/notesMasters/")) and n != "ppt/theme/theme2.xml"}
    master = ET.fromstring(parts["ppt/slideMasters/slideMaster1.xml"])
    master.find("p:cSld", NS).set("name", "Teal Editorial — shared branding")
    tree = master.find("p:cSld/p:spTree", NS)
    sample = ET.fromstring(data["ppt/slides/slide2.xml"])
    # Shared branding and a dynamic slide-number field live on the master.
    for sid in [2, 3, 4, 17, 18, 19]:
        shape = deepcopy(sample.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr[@id="{sid}"]]', namespaces=NS)[0])
        if sid == 17:
            replace_text(shape, "Presentation title • Presenter / Organisation")
            set_geometry(shape, w=15.5)
            shape.find("p:nvSpPr/p:cNvPr", NS).set("name", "EDIT HERE: presentation footer")
        if sid == 19:
            replace_text(shape, "1")
            shape.find("p:nvSpPr/p:cNvPr", NS).set("name", "Automatic slide number — do not replace with plain text")
            run = shape.find("p:txBody/a:p/a:r", NS)
            run.tag = q("a", "fld")
            run.set("id", "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "Teal Editorial slide number")) + "}")
            run.set("type", "slidenum")
        tree.append(shape)
    theme_refs(master)
    layoutlist = master.find("p:sldLayoutIdLst", NS)
    layoutlist.clear()
    master.find("p:hf", NS).set("sldNum", "1")
    theme = ET.fromstring(parts["ppt/theme/theme1.xml"])
    theme.set("name", "Teal Editorial")
    scheme = theme.find("a:themeElements/a:clrScheme", NS)
    scheme.set("name", "Teal Editorial")
    for slot, value in PALETTE.items():
        node = scheme.find(f"a:{slot}", NS)
        for child in list(node):
            node.remove(child)
        ET.SubElement(node, q("a", "srgbClr"), val=value)
    fonts = theme.find("a:themeElements/a:fontScheme", NS)
    fonts.set("name", "Georgia Editorial")
    for kind in ["majorFont", "minorFont"]:
        fonts.find(f"a:{kind}/a:latin", NS).set("typeface", "Georgia")
    parts["ppt/theme/theme1.xml"] = xml(theme)
    for i, (source_number, name, fields) in enumerate(LAYOUTS, 1):
        original = ET.fromstring(data[f"ppt/slides/slide{source_number}.xml"])
        layout = ET.Element(q("p", "sldLayout"), nsmap={k: NS[k] for k in ["a", "r", "p"]}, type="cust", preserve="1", matchingName=name)
        common = deepcopy(original.find("p:cSld", NS))
        common.set("name", name)
        layout.append(common)
        if source_number == 1:
            layout.set("showMasterSp", "0")
        shape_tree = common.find("p:spTree", NS)
        field_index = 0
        for shape in list(shape_tree.findall("p:sp", NS)):
            sid = int(shape.find("p:nvSpPr/p:cNvPr", NS).get("id"))
            text = texts(shape)
            if source_number != 1 and (sid in [2, 3, 4] or text.startswith("Understanding Neurodivergence —") or (not text and int(shape.find("p:spPr/a:xfrm/a:off", NS).get("y")) > 9500000)):
                shape_tree.remove(shape)
                continue
            if source_number == 1 and sid == 13:
                shape_tree.remove(shape)  # Stray empty text box in the original.
                continue
            if text.isdigit():
                shape_tree.remove(shape)
                continue
            if not text:
                continue
            role, example = fields[field_index]
            field_index += 1
            kind = "ctrTitle" if source_number == 1 and role == "Presentation title" else "title" if role == "Slide title" else "body"
            idx = 0 if kind in ["title", "ctrTitle"] else field_index
            placeholder(shape, idx, kind, role)
            replace_text(shape, "Click to add " + role.lower())
            if role in ["Section label", "Presentation type", "Discussion label", "Case label"]:
                set_geometry(shape, w=7 if role != "Section label" else 12)
            if role == "Presenter":
                set_geometry(shape, w=10, h=0.45)
            if source_number == 3 and role.startswith("Card"):
                card = int(role.split()[1]) - 1
                set_geometry(shape, x=1.56 + card * 4.4375, y=4.48 if role.endswith("heading") else 5.6, w=3.55, h=0.85 if role.endswith("heading") else 1.65)
            if source_number in [5, 10] and role.startswith("Block"):
                set_geometry(shape, w=8.05, h=0.42 if role.endswith("heading") else 0.73)
            # Clamp minor source overhangs to the original content margin.
            transform = shape.find("p:spPr/a:xfrm", NS)
            off, extent = transform.find("a:off", NS), transform.find("a:ext", NS)
            if int(off.get("x")) + int(extent.get("cx")) > round(19.5 * 914400):
                extent.set("cx", str(round(19.5 * 914400) - int(off.get("x"))))
        assert field_index == len(fields), (name, field_index, len(fields))
        color_map = original.find("p:clrMapOvr", NS)
        if color_map is not None:
            layout.append(deepcopy(color_map))
        theme_refs(layout)
        for shape in shape_tree.findall("p:sp", NS):
            ph = shape.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is not None and ph.get("type") in ["title", "ctrTitle"]:
                for font in shape.findall(".//a:latin", NS):
                    font.set("typeface", "+mj-lt")
        parts[f"ppt/slideLayouts/slideLayout{i}.xml"] = xml(layout)
        parts[f"ppt/slideLayouts/_rels/slideLayout{i}.xml.rels"] = relationships([("rId1", "slideMaster", "../slideMasters/slideMaster1.xml")])
        ET.SubElement(layoutlist, q("p", "sldLayoutId"), {"id": str(2147483648 + i), q("r", "id"): f"rId{i}"})
        slide = ET.Element(q("p", "sld"), nsmap={k: NS[k] for k in ["a", "r", "p"]})
        cs = ET.SubElement(slide, q("p", "cSld"), name=name)
        st = ET.SubElement(cs, q("p", "spTree"))
        for child in list(shape_tree)[:2]:
            st.append(deepcopy(child))
        by_role = dict(fields)
        for shape in shape_tree.findall("p:sp", NS):
            ph = shape.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is None:
                continue
            cloned = deepcopy(shape)
            cloned.find("p:nvSpPr/p:nvPr/p:ph", NS).attrib.pop("hasCustomPrompt", None)
            role = cloned.find("p:nvSpPr/p:cNvPr", NS).get("name")
            replace_text(cloned, by_role[role])
            st.append(cloned)
        override = ET.SubElement(slide, q("p", "clrMapOvr"))
        ET.SubElement(override, q("a", "masterClrMapping"))
        parts[f"ppt/slides/slide{i}.xml"] = xml(slide)
        parts[f"ppt/slides/_rels/slide{i}.xml.rels"] = relationships([("rId1", "slideLayout", f"../slideLayouts/slideLayout{i}.xml")])
    parts["ppt/slideMasters/slideMaster1.xml"] = xml(master)
    parts["ppt/slideMasters/_rels/slideMaster1.xml.rels"] = relationships([
        *[(f"rId{i}", "slideLayout", f"../slideLayouts/slideLayout{i}.xml") for i in range(1, 11)],
        ("rId11", "theme", "../theme/theme1.xml"),
    ])
    presentation = ET.fromstring(parts["ppt/presentation.xml"])
    notes = presentation.find("p:notesMasterIdLst", NS)
    if notes is not None:
        presentation.remove(notes)
    ids = presentation.find("p:sldIdLst", NS)
    ids.clear()
    for i in range(1, 11):
        ET.SubElement(ids, q("p", "sldId"), {"id": str(255 + i), q("r", "id"): f"rId{i + 1}"})
    parts["ppt/presentation.xml"] = xml(presentation)
    rels = ET.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    keep = []
    for rel in rels:
        kind = rel.get("Type").rsplit("/", 1)[-1]
        if kind not in ["slide", "notesMaster", "slideMaster"]:
            keep.append((rel.get("Id"), kind, rel.get("Target")))
    parts["ppt/_rels/presentation.xml.rels"] = relationships([
        ("rId1", "slideMaster", "slideMasters/slideMaster1.xml"),
        *[(f"rId{i + 1}", "slide", f"slides/slide{i}.xml") for i in range(1, 11)], *keep,
    ])
    # Regenerate overrides so removed notes/slides cannot leave dangling entries.
    types = ET.fromstring(parts["[Content_Types].xml"])
    for child in list(types):
        if child.tag == q("ct", "Override") and child.get("PartName").lstrip("/") not in parts:
            types.remove(child)
    existing = {c.get("PartName") for c in types}
    for i in range(1, 11):
        for folder, stem, content in [("slideLayouts", "slideLayout", "slideLayout"), ("slides", "slide", "slide")]:
            name = f"/ppt/{folder}/{stem}{i}.xml"
            if name not in existing:
                ET.SubElement(types, q("ct", "Override"), PartName=name, ContentType=f"application/vnd.openxmlformats-officedocument.presentationml.{content}+xml")
    # Remove the old deck's identifying metadata, notes counts and title list.
    for name in ["docProps/core.xml", "docProps/app.xml"]:
        root = ET.fromstring(parts[name])
        for child in list(root):
            root.remove(child)
        if name.endswith("core.xml"):
            ET.SubElement(root, "{http://purl.org/dc/elements/1.1/}title").text = "Teal Editorial Presentation Template"
        else:
            ET.SubElement(root, "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Application").text = "Microsoft Office PowerPoint"
            ET.SubElement(root, "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Slides").text = "10"
        parts[name] = xml(root)
    for extension, kind in [("pptx", "presentation"), ("potx", "template")]:
        types.find('ct:Override[@PartName="/ppt/presentation.xml"]', NS).set("ContentType", f"application/vnd.openxmlformats-officedocument.presentationml.{kind}.main+xml")
        parts["[Content_Types].xml"] = xml(types)
        path = ROOT / f"Teal_Editorial_{'Starter' if extension == 'pptx' else 'Template'}.{extension}"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        validate(path)
        print(f"Created and validated {path.name}")


def validate(path):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                root = ET.fromstring(archive.read(name))
                if name.endswith(".rels"):
                    base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        if rel.get("TargetMode") != "External":
                            target = posixpath.normpath(posixpath.join(base, rel.get("Target")))
                            assert target in archive.namelist(), (name, target)
        for i in range(1, 11):
            root = ET.fromstring(archive.read(f"ppt/slideLayouts/slideLayout{i}.xml"))
            phs = root.findall(".//p:ph", NS)
            assert len(phs) == len(LAYOUTS[i - 1][2])
            assert len({p.get("idx") for p in phs}) == len(phs)
            assert len(root.xpath('.//p:ph[@type="title" or @type="ctrTitle"]', namespaces=NS)) == 1
            assert not any("Neurodivergence" in t for t in root.xpath(".//a:t/text()", namespaces=NS))
            for shape in root.findall(".//p:sp", NS):
                transform = shape.find("p:spPr/a:xfrm", NS)
                off, ext = transform.find("a:off", NS), transform.find("a:ext", NS)
                assert int(off.get("x")) >= 0 and int(off.get("y")) >= 0
                # Cover's original background has tiny rounding overhangs.
                assert int(off.get("x")) + int(ext.get("cx")) <= 18288000 + 500
                assert int(off.get("y")) + int(ext.get("cy")) <= 10287000 + 500
        # python-pptx accepts presentation parts, not POTX: change only the MIME
        # in memory for this independent parser check, leaving the POTX intact.
        payload = BytesIO()
        with zipfile.ZipFile(payload, "w") as copy:
            for name in archive.namelist():
                value = archive.read(name)
                if name == "[Content_Types].xml":
                    value = value.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                copy.writestr(name, value)
    payload.seek(0)
    deck = Presentation(payload)
    assert len(deck.slides) == len(deck.slide_layouts) == 10
    for i, layout in enumerate(deck.slide_layouts):
        assert layout.name == LAYOUTS[i][1]
        fresh = deck.slides.add_slide(layout)
        assert fresh.shapes.title is not None
        assert len(fresh.placeholders) == len(LAYOUTS[i][2])
        for ph in fresh.placeholders:
            assert ph.width > 0 and ph.height > 0
            assert ph.placeholder_format.idx in {p.placeholder_format.idx for p in layout.placeholders}
    # Exercise saving/reopening newly inserted slides, not just the specimens.
    check = BytesIO()
    deck.save(check)
    check.seek(0)
    assert len(Presentation(check).slides) == 20


if __name__ == "__main__":
    build()
